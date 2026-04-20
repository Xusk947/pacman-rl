from __future__ import annotations

from pathlib import Path
from typing import Any


def write_telemetry_xlsx(path: Path, *, rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
    except ModuleNotFoundError as e:
        raise ModuleNotFoundError("openpyxl is required to export telemetry to .xlsx") from e

    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "metrics"

    if not rows:
        ws["A1"] = "no data"
        wb.save(path)
        return

    keys = sorted({k for r in rows for k in r.keys()})
    ws.append(keys)
    for r in rows:
        ws.append([r.get(k) for k in keys])

    def add_chart(title: str, y_keys: list[str], anchor: str) -> None:
        existing = [k for k in y_keys if k in keys]
        if not existing:
            return

        x_idx = keys.index("update") + 1 if "update" in keys else 1

        chart = LineChart()
        chart.title = title
        chart.y_axis.title = "value"
        chart.x_axis.title = "update"

        min_row = 1
        max_row = 1 + len(rows)
        cats = Reference(ws, min_col=x_idx, min_row=min_row + 1, max_row=max_row)
        chart.set_categories(cats)

        for k in existing:
            col = keys.index(k) + 1
            data = Reference(ws, min_col=col, min_row=min_row, max_row=max_row)
            chart.add_data(data, titles_from_data=True)

        ws.add_chart(chart, anchor)

    add_chart(
        "Rewards",
        ["pacman_reward_mean", "ghosts_reward_mean"],
        "A" + str(len(rows) + 3),
    )
    add_chart(
        "Loss / KL",
        ["pacman_loss", "pacman_approx_kl", "ghosts_loss", "ghosts_approx_kl"],
        "I" + str(len(rows) + 3),
    )

    wb.save(path)

